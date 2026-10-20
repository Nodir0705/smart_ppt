"""Generalization regression net over the generated deck corpus.

`tools/make_corpus.build_corpus` emits small decks whose pointable texts are
known by construction. Here we parse each and assert no author-intended element
is silently dropped — the failure mode that hides when everything is tuned to
one deck.

Decks that target a *known* parser gap (e.g. grouped shapes) assert the gap
still exists, so when the parser is later fixed this test fails loudly and tells
us to retire the `known_gap` marker — a self-updating regression net.
"""

from __future__ import annotations

import pytest

from src.group_scope import expansions_for, scope_index
from src.parser import _clean_alt_text, parse_deck
from tools.make_corpus import build_corpus


@pytest.mark.parametrize("raw,expected", [
    ("Architecture diagram", "Architecture diagram"),  # real alt text kept
    ("Logo", "Logo"),
    ("/home/x/photo.png", ""),                          # file path rejected
    ("C:\\img\\a.PNG", ""),
    ("photo.jpg", ""),                                   # bare filename rejected
    ("Picture 1", ""),                                   # auto-name rejected
    ("image", ""),
    ("  ", ""),
    # PowerPoint's auto-generated description: a generic caption whose words
    # the slide never contains, always tagged with the marker sentence
    ("A close-up of a graph\n\nAI-generated content may be incorrect.", ""),
    ("A diagram of a machine\n\nai-generated content may be incorrect", ""),
    ("AI-Generated Content May Be Incorrect.", ""),
])
def test_clean_alt_text_rejects_junk(raw, expected):
    assert _clean_alt_text(raw) == expected


def _two_paragraph_box():
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tf = tb.text_frame
    tf.text = "Instructor:  Mike Sipser"
    tf.add_paragraph().text = "Second paragraph"
    return tb


def test_plain_paragraph_text_unchanged():
    # Characterisation: ordinary runs stay verbatim (internal double space
    # included) and paragraphs join with "\n" — pinned so the math/br-aware
    # paragraph reader cannot drift for the shapes that have neither.
    from src.parser import _paragraph_metrics, _paragraphs, _shape_text
    tb = _two_paragraph_box()
    assert _shape_text(tb) == "Instructor:  Mike Sipser\nSecond paragraph"
    assert _paragraphs(tb) == ["Instructor:  Mike Sipser", "Second paragraph"]
    assert [m["text"] for m in _paragraph_metrics(tb)] == _paragraphs(tb)


def test_soft_line_break_paragraph_spans_two_lines():
    # paragraph 1 = "a" + br + "b" occupies TWO lines, so its box is two line
    # heights tall and paragraph 2 starts two line heights below its top — not
    # one, which drifted every bullet after a br-paragraph upward onto it.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from src.parser import _bullet_bboxes
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    p1 = tb.text_frame.paragraphs[0]
    p1.add_run().text = "a"
    p1.add_line_break()
    p1.add_run().text = "b"
    p1.font.size = Pt(20)
    p2 = tb.text_frame.add_paragraph()
    p2.text = "c"
    p2.font.size = Pt(20)

    (t1, b1), (t2, b2) = _bullet_bboxes(tb)
    assert (t1, t2) == ("a\nb", "c")
    one_line = Pt(20).inches * 1.2  # 20 pt at the default 1.2 line spacing
    assert abs(b1.h - 2 * one_line) < 1e-6, b1
    assert abs(b2.y - (b1.y + 2 * one_line)) < 1e-6, (b1, b2)


def _a_p(inner: str):
    from lxml import etree
    return etree.fromstring(
        '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"'
        f' xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{inner}</a:p>')


def test_paragraph_text_reads_math_runs_and_keeps_line_breaks_as_newlines():
    # Equation runs sit in <a14:m> between ordinary runs (python-pptx skips them);
    # their math-italic letters normalise to plain ones and U+00A0 to a space.
    # A soft line break (<a:br/>) is a line boundary — "\n", the same separator
    # paragraphs get — so each br-separated line stays its own trigger phrase.
    # Ordinary runs are verbatim: no whitespace is absorbed around the break.
    from src.parser import _paragraph_text
    math = ("<a14:m><m:oMath><m:r><m:t>\U0001D434</m:t></m:r><m:r><m:t>, </m:t></m:r>"
            "<m:r><m:t>\U0001D435</m:t></m:r></m:oMath></a14:m>")
    p = _a_p(f"<a:r><a:t>Let </a:t></a:r>{math}<a:r><a:t> be languages</a:t></a:r>")
    assert _paragraph_text(p) == "Let A, B be languages"
    p = _a_p("<a:r><a:t>TAs:</a:t></a:r><a:br/><a:r><a:t>Fadi</a:t></a:r>")
    assert _paragraph_text(p) == "TAs:\nFadi"
    p = _a_p("<a:r><a:t>Examples: </a:t></a:r><a:br/><a:br/><a:r><a:t> program</a:t></a:r>")
    assert _paragraph_text(p) == "Examples: \n\n program"


def _element_texts(index: dict) -> list[str]:
    return [
        (e.get("text") or "").strip()
        for s in index["slides"]
        for e in s["elements"]
    ]


def _survives(needle: str, haystack: list[str]) -> bool:
    n = needle.lower()
    return any(n in h.lower() for h in haystack)


def test_corpus_no_silent_drops(tmp_path):
    decks = build_corpus(tmp_path)
    assert decks, "corpus generator produced no decks"

    for d in decks:
        if getattr(d, "requires_ocr", False):
            continue  # pure parse can't OCR; the slow OCR test covers these
        index = parse_deck(d.path)  # must not raise (covers empty/edge decks)
        texts = _element_texts(index)
        n_elements = len(texts)
        present = [t for t in d.expected_texts if _survives(t, texts)]

        if d.known_gap:
            # The deck targets a documented gap: its texts SHOULD survive but
            # don't yet. If they start surviving, the gap is fixed — fail so we
            # remove the marker rather than silently keep xfailing forever.
            assert len(present) < len(d.expected_texts), (
                f"{d.name}: known gap appears FIXED ({d.known_gap}). "
                f"Retire known_gap in tools/make_corpus.py and assert survival."
            )
            continue  # element-count expectations describe the post-fix deck

        missing = [t for t in d.expected_texts if not _survives(t, texts)]
        assert not missing, f"{d.name}: parser silently dropped {missing}"
        assert n_elements >= d.expected_min_elements, (
            f"{d.name}: expected >= {d.expected_min_elements} elements, got {n_elements}"
        )

        if d.expected_types:
            from collections import Counter
            got = Counter(e["type"] for s in index["slides"] for e in s["elements"])
            need = Counter(d.expected_types)
            for typ, count in need.items():
                assert got[typ] >= count, (
                    f"{d.name}: expected >= {count} '{typ}' elements, got {got[typ]}"
                )


def test_rotated_shape_bbox_is_aabb(tmp_path):
    # a 6x1in box rotated 90deg -> AABB should swap to ~1x6 (tall, narrow)
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["rotated_shape"].path)
    el = next(e for s in index["slides"] for e in s["elements"]
              if (e.get("text") or "").startswith("Rotated"))
    b = el["bbox"]
    assert b["h"] > b["w"], f"rotation not applied: w={b['w']:.2f} h={b['h']:.2f}"
    assert abs(b["w"] - 1.0) < 0.05 and abs(b["h"] - 6.0) < 0.05, b


def test_parse_deck_isolates_a_failing_slide(tmp_path, monkeypatch, caplog):
    # A malformed slide must warn + yield an empty slide, not abort the deck.
    import src.parser as P
    deck = {d.name: d for d in build_corpus(tmp_path)}["title_and_bullets"]
    real = P._parse_slide

    def boom(slide, i, w, h, *, ocr=False):
        if i == 1:
            raise ValueError("synthetic parse failure")
        return real(slide, i, w, h, ocr=ocr)

    monkeypatch.setattr(P, "_parse_slide", boom)
    with caplog.at_level("WARNING"):
        index = P.parse_deck(deck.path)  # must NOT raise

    assert index["slides"], "deck aborted instead of isolating the bad slide"
    assert index["slides"][0]["slide_number"] == 1
    assert index["slides"][0]["elements"] == []  # failed slide → empty, numbered
    assert any("failed to parse" in r.message for r in caplog.records)


def test_unsupported_chart_is_needs_vision_not_fabricated(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["radar_chart"].path)
    els = [e for s in index["slides"] for e in s["elements"]]
    charts = [e for e in els if e["type"] == "chart"]
    cats = [e for e in els if e["type"] == "chart_category"]
    assert len(charts) == 1 and charts[0]["needs_vision"] is True
    assert cats == [], f"radar fabricated {len(cats)} category bars"


def test_grouped_chart_has_distinct_subcolumns(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["grouped_bar"].path)
    els = [e for s in index["slides"] for e in s["elements"]]
    chart = next(e for e in els if e["type"] == "chart")
    cats = [e for e in els if e["type"] == "chart_category"]

    assert chart["multi_series"] is True
    assert len(cats) == 6, "expected 3 categories x 2 series"
    # distinct, non-overlapping sub-columns in category-major x order
    xs = [c["bbox"]["x"] for c in cats]
    assert xs == sorted(xs) and len(set(xs)) == 6
    for a, b in zip(cats, cats[1:]):
        assert a["bbox"]["x"] + a["bbox"]["w"] <= b["bbox"]["x"] + 1e-6, "sub-columns overlap"
    # value-proportional height: Revenue 30 (tallest) > Cost 5 (shortest)
    h = {(c["text"]): c["bbox"]["h"] for c in cats}
    assert h["Q3 — Revenue"] > h["Q1 — Cost"]
    # series name reaches the triggers (disambiguation handle)
    rev = next(c for c in cats if c["series"] == "Revenue" and "Q2" in c["text"])
    assert any("revenue" in p for p in rev["trigger_phrases"])


def test_stacked_chart_segments_stack_within_category(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["stacked_bar"].path)
    els = [e for s in index["slides"] for e in s["elements"]]
    assert next(e for e in els if e["type"] == "chart")["multi_series"] is True
    cats = [e for e in els if e["type"] == "chart_category"]
    assert len(cats) == 6

    # within a category the two segments share x and stack contiguously (the
    # upper series' bottom == the lower series' top), no overlap
    q1 = [c for c in cats if c["text"].startswith("Q1")]
    base, bonus = q1[0]["bbox"], q1[1]["bbox"]   # category-major: Base then Bonus
    assert abs(base["x"] - bonus["x"]) < 1e-6
    assert base["y"] > bonus["y"]                 # Base lower on screen than Bonus
    assert abs((bonus["y"] + bonus["h"]) - base["y"]) < 1e-6, "segments not contiguous"
    # taller stack for the larger category total (Q3 total 42 > Q1 total 15)
    q3_base = next(c for c in cats if c["text"] == "Q3 — Base")
    assert q3_base["bbox"]["h"] > base["h"]


def test_multiseries_scoped_phrases_are_speakable(tmp_path):
    # scope_index builds chart_category phrases from the element label; the
    # multi-series label is "Cat — Series", which must not leak the em-dash (STT
    # never says it) or produce "... bar chart chart".
    from src.group_scope import scope_index
    decks = {d.name: d for d in build_corpus(tmp_path)}
    for name in ("grouped_bar", "stacked_bar"):
        index = parse_deck(decks[name].path)
        scope_index(index)
        phrases = [p for s in index["slides"] for e in s["elements"]
                   if e["type"] == "chart_category" for p in e["trigger_phrases"]]
        assert phrases, f"{name}: no scoped phrases"
        assert not [p for p in phrases if "—" in p], f"{name}: em-dash leaked into a phrase"
        assert not [p for p in phrases if "chart chart" in p], f"{name}: double 'chart'"
        # the series name still reaches a phrase for disambiguation
        assert any("revenue" in p or "base" in p for p in phrases)


def test_pie_has_distinct_value_proportional_wedges(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["pie_chart"].path)
    els = [e for s in index["slides"] for e in s["elements"]]
    chart = next(e for e in els if e["type"] == "chart")
    cats = [e for e in els if e["type"] == "chart_category"]
    # pie has sound geometry -> computed wedges, not a needs_vision punt
    assert chart["needs_vision"] is False
    assert len(cats) == 4
    # the four wedge AABBs are distinct (each slice is a different region)
    boxes = {(round(c["bbox"]["x"], 3), round(c["bbox"]["y"], 3),
              round(c["bbox"]["w"], 3), round(c["bbox"]["h"], 3)) for c in cats}
    assert len(boxes) == 4, "pie wedges collapsed to identical boxes"


def test_nested_groups_map_into_slide_space(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["nested_groups"].path)
    els = {(e.get("text") or ""): e for s in index["slides"] for e in s["elements"]}
    # depth-2 child: identity child-extents -> stays at its authored slide coords
    deep = els["Deep node"]["bbox"]
    assert abs(deep["x"] - 2.0) < 0.05 and abs(deep["y"] - 2.0) < 0.05, deep
    mid = els["Mid node"]["bbox"]
    assert abs(mid["x"] - 5.0) < 0.05 and abs(mid["y"] - 3.0) < 0.05, mid


def test_duplicate_text_not_deduped(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["duplicate_text"].path)
    live = [e for s in index["slides"] for e in s["elements"]
            if (e.get("text") or "") == "Live"]
    assert len(live) == 2, "identical-text boxes were collapsed"
    xs = sorted(round(e["bbox"]["x"], 1) for e in live)
    assert xs == [1.0, 6.0], xs


@pytest.mark.slow
@pytest.mark.skipif(__import__("shutil").which("soffice") is None,
                    reason="soffice required to render the PDF text layer")
def test_wrapping_paragraph_refines_to_stacked_line_bboxes(tmp_path):
    """The refine step must split ONE soft-wrapping paragraph into multiple
    per-line boxes that stack top-to-bottom without overlap and union back to
    the element box — the contract line-level highlighting consumes.

    Structural asserts only (count >= 2, monotonic y, left-aligned, contiguous
    union) so it survives soffice font/version differences in exact wrapping.
    """
    from src.refine_text import refine_index
    from src.render import deck_pdf

    decks = {d.name: d for d in build_corpus(tmp_path)}
    deck = decks["wrapping_paragraph"]
    index = parse_deck(deck.path)
    # one element pre-refine (no paragraph breaks to split on)
    els = [e for s in index["slides"] for e in s["elements"]]
    assert len(els) == 1, f"wrap paragraph split into {len(els)} elements pre-refine"

    pdf = deck_pdf(deck.path, tmp_path)
    refine_index(index, pdf)
    el = [e for s in index["slides"] for e in s["elements"]][0]
    lbs = el["line_bboxes"]

    assert el["bbox_source"] == "pdf_text"
    assert len(lbs) >= 2, f"expected a multi-line wrap, got {len(lbs)} line(s)"
    # stacked top-to-bottom, no vertical overlap between consecutive lines
    ys = [b["y"] for b in lbs]
    assert ys == sorted(ys), "line_bboxes not ordered top-to-bottom"
    for a, b in zip(lbs, lbs[1:]):
        assert a["y"] + a["h"] <= b["y"] + 1e-6, "consecutive lines overlap vertically"
    # left-aligned: all lines start within a small x band
    xs = [b["x"] for b in lbs]
    assert max(xs) - min(xs) < 0.1, f"lines not left-aligned: {xs}"
    # element bbox is the union of its line boxes
    ux = min(xs)
    uy = min(ys)
    ux1 = max(b["x"] + b["w"] for b in lbs)
    uy1 = max(b["y"] + b["h"] for b in lbs)
    bb = el["bbox"]
    assert abs(bb["x"] - ux) < 1e-3 and abs(bb["y"] - uy) < 1e-3
    assert abs(bb["w"] - (ux1 - ux)) < 1e-3 and abs(bb["h"] - (uy1 - uy)) < 1e-3

    # line_texts must reconstruct CLEAN words (not char-fragmented glyph chunks),
    # in order, and concatenate back to the original paragraph — this is the
    # handle select_line matches the spoken tail against.
    lts = el["line_texts"]
    assert len(lts) == len(lbs), "line_texts not parallel to line_bboxes"
    para = el["text"].split()
    assert " ".join(lts).split() == para, f"line_texts != paragraph words: {lts}"
    # at least one line carries multiple whole words (guards the char-fragment bug)
    assert any(len(lt.split()) >= 2 for lt in lts), f"lines look fragmented: {lts}"


@pytest.mark.slow
@pytest.mark.skipif(__import__("shutil").which("tesseract") is None,
                    reason="tesseract required for the OCR fallback")
def test_ocr_recovers_baked_image_text(tmp_path):
    # An alt-text-less image with baked-in text is unreachable by pure parsing;
    # ocr=True must recover it and clear needs_vision.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    deck = decks["image_baked_text"]
    plain = parse_deck(deck.path)              # no OCR -> still needs_vision
    img0 = [e for s in plain["slides"] for e in s["elements"] if e["type"] == "image"][0]
    assert img0["needs_vision"] is True and not (img0.get("text") or "")

    index = parse_deck(deck.path, ocr=True)
    imgs = [e for s in index["slides"] for e in s["elements"] if e["type"] == "image"]
    assert len(imgs) == 1
    assert "NET REVENUE" in (imgs[0].get("text") or "").upper()
    assert imgs[0]["needs_vision"] is False
    assert imgs[0]["trigger_phrases"], "OCR text produced no trigger phrases"


@pytest.mark.slow
@pytest.mark.skipif(__import__("shutil").which("tesseract") is None,
                    reason="tesseract required for the OCR fallback")
def test_ocr_does_not_hallucinate_on_textless_image(tmp_path):
    # Precision guard: shapes-only image must NOT yield a pointable, even with OCR
    # on (the conf + min-alpha filters must suppress tesseract's hallucinations).
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["image_no_text"].path, ocr=True)
    imgs = [e for s in index["slides"] for e in s["elements"] if e["type"] == "image"]
    assert len(imgs) == 1
    assert not (imgs[0].get("text") or ""), f"OCR hallucinated: {imgs[0].get('text')!r}"
    assert imgs[0]["needs_vision"] is True
    assert imgs[0]["trigger_phrases"] == []


def _texts(index: dict) -> list[str]:
    return [e["text"] or "" for s in index["slides"] for e in s["elements"]]


def test_hidden_shapes_are_not_indexed(tmp_path):
    # `cNvPr hidden="1"` shapes are invisible on the rendered slide; a hidden
    # group hides its whole subtree. The visible sibling is untouched.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    texts = _texts(parse_deck(decks["hidden_shapes"].path))
    assert texts == ["Visible sibling"], texts


def test_off_canvas_shapes_are_not_indexed(tmp_path):
    # Wholly outside the slide (y=9 in on a 7.5-in slide) → invisible → dropped;
    # a box overhanging the bottom edge (y=7.3, h=0.5) is still partly visible.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    texts = _texts(parse_deck(decks["off_canvas"].path))
    assert texts == ["On the slide", "Overhanging the edge"], texts


def test_text_under_a_later_opaque_shape_is_not_indexed(tmp_path):
    # z-order = document order: a later opaque rect fully covering a text box
    # hides it on the rendered slide. Controls (one slide each) stay indexed:
    # noFill cover, 50 % alpha cover, rect drawn before the text, partial cover,
    # an ANIMATED cover (entrance/exit effect: the audience sees what is under
    # it at some point), and shapes whose AABB exceeds their paint — a rotated
    # bar, an ellipse — with the text in the unpainted part of the AABB.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["occluded_text"].path)
    per_slide = [[e["text"] for e in s["elements"]] for s in index["slides"]]
    assert per_slide == [
        [],
        ["Cover has no fill"],
        ["Cover is half transparent"],
        ["Rect is underneath"],
        ["Only partly covered"],
        ["Cover enters on click"],
        ["Beside a rotated bar"],
        ["In an ellipse's corner"],
    ], per_slide


@pytest.mark.parametrize("deck,bbox", [
    ("placeholder_picture", None),                 # the layout's placeholder geometry
    ("movie", {"x": 2.0, "y": 1.5, "w": 6.0, "h": 4.0}),
])
def test_every_p_pic_is_an_image(tmp_path, deck, bbox):
    # A picture in a layout placeholder (PLACEHOLDER) and an embedded video
    # (MEDIA) are both `p:pic` elements → exactly one needs_vision image each.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    els = [e for s in parse_deck(decks[deck].path)["slides"] for e in s["elements"]]
    assert [e["type"] for e in els] == ["image"], els
    assert els[0]["needs_vision"] is True and els[0]["text"] is None
    if bbox:
        assert {k: round(els[0]["bbox"][k], 2) for k in bbox} == bbox
    # OCR reads a movie's poster frame (no `.image`) — must not raise either way
    assert [e["type"] for s in parse_deck(decks[deck].path, ocr=True)["slides"]
            for e in s["elements"]] == ["image"]


def test_arrows_come_from_geometry_not_from_being_small(tmp_path):
    # connector, RIGHT_ARROW preset and a thin filled 3x0.05 bar are arrows;
    # the 0.1-in dot and the empty zero-width text box are decorative — absent.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    els = [e for s in parse_deck(decks["arrow_shapes"].path)["slides"] for e in s["elements"]]
    assert [e["type"] for e in els] == ["arrow"] * 3, els
    assert sorted((round(e["bbox"]["w"], 2), round(e["bbox"]["h"], 2)) for e in els) == [
        (2.0, 0.5), (3.0, 0.05), (3.0, 0.5)]


def test_box_holding_several_texts_is_not_a_label_box(tmp_path):
    # A full-slide empty rect behind a title + two body boxes is a background:
    # it pairs with nothing (no diagram_label/stat with its bbox); the title is
    # typed by `_classify_text_shape` (y < 1.2 in, wide) with its own bbox.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    els = [e for s in parse_deck(decks["background_panel"].path)["slides"] for e in s["elements"]]
    assert [(e["type"], e["text"]) for e in els] == [
        ("title", "Panel Title"),
        ("body", "First point on the panel"),
        ("body", "Second point on the panel"),
    ], els
    assert round(els[0]["bbox"]["w"], 2) == 8.0 and round(els[0]["bbox"]["y"], 2) == 0.4


def test_ink_annotation_resolves_to_its_fallback_picture(tmp_path):
    # PowerPoint ink: an mc:AlternateContent whose p14 Choice is a
    # `p:contentPart` (python-pptx has no class for it — it used to crash the
    # slide, which was then emitted EMPTY) and whose Fallback is a picture of
    # the strokes. The title survives; the ink is one needs_vision image.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    els = [e for s in parse_deck(decks["ink_annotation"].path)["slides"] for e in s["elements"]]
    assert [(e["type"], e["text"]) for e in els] == [
        ("title", "Visible title text"), ("image", None)], els
    assert els[1]["needs_vision"] is True


def test_table_cell_reads_equation_runs(tmp_path):
    # Cell text goes through the same paragraph reader as text boxes, so an
    # `a14:m` equation in a cell (q₁: math-italic q over a subscript 1) reads
    # "q1" instead of python-pptx's empty `cell.text`; plain cells unchanged.
    decks = {d.name: d for d in build_corpus(tmp_path)}
    els = [e for s in parse_deck(decks["table_math_cell"].path)["slides"] for e in s["elements"]]
    assert [e["text"] for e in els] == ["δ", "0", "start", "q1"], els


def test_corpus_parse_is_deterministic(tmp_path):
    # The generator + parser are deterministic: same build → same parsed texts.
    a = build_corpus(tmp_path / "a")
    b = build_corpus(tmp_path / "b")
    by_name_b = {d.name: d for d in b}
    for da in a:
        ta = _element_texts(parse_deck(da.path))
        tb = _element_texts(parse_deck(by_name_b[da.name].path))
        assert ta == tb, f"{da.name}: non-deterministic parse {ta!r} != {tb!r}"


def test_chart_title_and_series_names_emitted(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["titled_chart_acronym"].path)
    chart = next(e for s in index["slides"] for e in s["elements"] if e["type"] == "chart")
    assert chart["chart_title"] == "Revenue by region (USD)"
    assert chart["series_names"] == ["Revenue"]
    # chart text is still the chart-type token (container embedding unchanged)
    assert chart["text"].startswith("COLUMN_CLUSTERED")


def test_untitled_chart_has_no_chart_title(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["pie_chart"].path)
    chart = next(e for s in index["slides"] for e in s["elements"] if e["type"] == "chart")
    assert chart.get("chart_title") is None
    assert chart["series_names"] == ["Share"]


def test_multiseries_chart_lists_every_series(tmp_path):
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["grouped_bar"].path)
    chart = next(e for s in index["slides"] for e in s["elements"] if e["type"] == "chart")
    assert len(chart["series_names"]) >= 2


def test_titled_chart_acronym_scopes_without_deck_vocab(tmp_path):
    """Group label comes from the chart's own title and STT's expansion from the
    deck's own 'Speech to Text (STT)' — no test_deck vocabulary involved."""
    decks = {d.name: d for d in build_corpus(tmp_path)}
    index = parse_deck(decks["titled_chart_acronym"].path)
    out = scope_index(index, expansions=expansions_for(index, str(tmp_path / "none.json")))
    els = {e.get("text"): e for s in out["slides"] for e in s["elements"]}
    north = els["North"]
    assert north["group_label"] == "revenue"
    assert "the north bar in the revenue chart" in north["trigger_phrases"]
    assert "the north bar in the revenue by region chart" in north["trigger_phrases"]
    engine = els["STT Engine"]["trigger_phrases"]
    assert "speech to text" in engine
    assert "speech to text engine" in engine


@pytest.mark.skipif(__import__("shutil").which("tesseract") is None,
                    reason="tesseract required for the OCR fallback")
def test_ocr_preserves_element_ids(tmp_path):
    """The bench narrates from the ocr=False truth parse and scores the ocr=True
    index — that only works if OCR fills text on existing elements, never adds,
    drops or renumbers them."""
    for d in build_corpus(tmp_path):
        ids = lambda ix: [e["id"] for s in ix["slides"] for e in s["elements"]]
        assert ids(parse_deck(d.path, ocr=False)) == ids(parse_deck(d.path, ocr=True)), d.name


def test_build_index_accepts_ocr_flag():
    import inspect
    from src.build_index import build_index
    assert inspect.signature(build_index).parameters["ocr"].default is True
