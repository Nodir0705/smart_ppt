"""Tests for PDF caching (render.deck_pdf) and text-bbox refinement."""
from __future__ import annotations

import shutil
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "test_deck.pptx"
INDEX_JSON = ROOT / "out" / "index.json"

SOFFICE_MISSING = shutil.which("soffice") is None


@pytest.mark.skipif(SOFFICE_MISSING, reason="soffice not installed")
def test_deck_pdf_caches(tmp_path):
    from src.render import deck_pdf

    pdf = deck_pdf(DECK, tmp_path)
    assert pdf.exists()
    assert pdf.suffix == ".pdf"
    mtime = pdf.stat().st_mtime

    # Second call must be a cache hit: same path, not regenerated.
    pdf2 = deck_pdf(DECK, tmp_path)
    assert pdf2 == pdf
    assert pdf2.stat().st_mtime == mtime


def test_words_to_inches_scale_one():
    from src.refine_text import words_to_inches

    words = [{"text": "Hi", "x0": 72.0, "x1": 144.0, "top": 36.0, "bottom": 72.0}]
    out = words_to_inches(words, 1.0, 1.0)
    assert out[0]["text"] == "Hi"
    assert out[0]["x0"] == pytest.approx(1.0)
    assert out[0]["x1"] == pytest.approx(2.0)
    assert out[0]["y0"] == pytest.approx(0.5)
    assert out[0]["y1"] == pytest.approx(1.0)


def test_assign_words_smallest_box_wins():
    from src.refine_text import assign_words

    big = {"x": 0.0, "y": 0.0, "w": 10.0, "h": 10.0, "units": "inches"}
    small = {"x": 1.0, "y": 1.0, "w": 2.0, "h": 2.0, "units": "inches"}
    word = {"text": "x", "x0": 1.5, "y0": 1.5, "x1": 2.0, "y1": 2.0}  # center (1.75,1.75)
    assigned = assign_words([("big", big), ("small", small)], [word])
    assert assigned["small"] == [word]
    assert assigned["big"] == []


def test_assign_words_drops_outside():
    from src.refine_text import assign_words

    box = {"x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0, "units": "inches"}
    word = {"text": "x", "x0": 5.0, "y0": 5.0, "x1": 5.5, "y1": 5.5}
    assigned = assign_words([("e", box)], [word])
    assert assigned["e"] == []


def test_cluster_lines_two_rows():
    from src.refine_text import cluster_lines

    w = lambda t, x0, y0: {"text": t, "x0": x0, "y0": y0, "x1": x0 + 0.5, "y1": y0 + 0.2}
    words = [w("b", 2.0, 0.0), w("a", 1.0, 0.0), w("c", 1.0, 1.0)]  # row0: a,b ; row1: c
    lines = cluster_lines(words, y_tol=0.03)
    assert [[x["text"] for x in line] for line in lines] == [["a", "b"], ["c"]]


def test_union_bbox_tight():
    from src.refine_text import union_bbox

    words = [
        {"x0": 1.0, "y0": 1.0, "x1": 2.0, "y1": 1.5},
        {"x0": 2.5, "y0": 1.1, "x1": 3.0, "y1": 1.4},
    ]
    box = union_bbox(words)
    assert box == {"x": 1.0, "y": 1.0, "w": 2.0, "h": 0.5, "units": "inches"}


def test_apply_refinement_no_words_keeps_shape():
    from src.refine_text import _apply_refinement

    e = {"id": "e1", "type": "body", "text": "hi",
         "bbox": {"x": 1.0, "y": 1.0, "w": 4.0, "h": 2.0, "units": "inches"}}
    _apply_refinement(e, [])
    assert e["bbox_source"] == "shape"
    assert e["line_bboxes"] == []
    assert e["shape_bbox"] == {"x": 1.0, "y": 1.0, "w": 4.0, "h": 2.0, "units": "inches"}
    assert e["bbox"] == e["shape_bbox"]


def test_apply_refinement_sets_pdf_text():
    from src.refine_text import _apply_refinement

    e = {"id": "e1", "type": "body", "text": "hi there",
         "bbox": {"x": 0.0, "y": 0.0, "w": 9.0, "h": 3.0, "units": "inches"}}
    words = [
        {"text": "hi", "x0": 1.0, "y0": 1.0, "x1": 1.5, "y1": 1.2},
        {"text": "there", "x0": 1.6, "y0": 1.0, "x1": 2.4, "y1": 1.2},
    ]
    _apply_refinement(e, words)
    assert e["bbox_source"] == "pdf_text"
    assert e["shape_bbox"]["w"] == 9.0
    # x/w are the tight glyph union; pypdfium2 tight boxes are already the ink
    # extent (SHIFT_FRAC=0.0), so y is the glyph-union top unchanged; h unchanged.
    assert e["bbox"] == {"x": 1.0, "y": 1.0, "w": 1.4, "h": 0.2, "units": "inches"}
    assert len(e["line_bboxes"]) == 1
    assert e["bbox"]["w"] < e["shape_bbox"]["w"]


def test_cluster_lines_no_cascade_on_sliver_overlap():
    from src.refine_text import cluster_lines

    # consecutive line boxes overlap by only 0.05 (tight leading) -> 3 distinct lines
    w = lambda t, y0: {"text": t, "x0": 1.0, "y0": y0, "x1": 1.5, "y1": y0 + 0.2}
    words = [w("r1", 1.0), w("r2", 1.15), w("r3", 1.30)]
    lines = cluster_lines(words, y_tol=0.03)
    assert [[x["text"] for x in ln] for ln in lines] == [["r1"], ["r2"], ["r3"]]


def test_cluster_lines_groups_within_y_tol():
    from src.refine_text import cluster_lines

    # two short words, boxes do not overlap but centers are within y_tol -> same line
    a = {"text": "a", "x0": 1.0, "y0": 1.000, "x1": 1.4, "y1": 1.010}
    b = {"text": "b", "x0": 1.5, "y0": 1.011, "x1": 1.9, "y1": 1.021}
    lines = cluster_lines([a, b], y_tol=0.03)
    assert [[x["text"] for x in ln] for ln in lines] == [["a", "b"]]


def _slide(index, n):
    return next(s for s in index["slides"] if s["slide_number"] == n)


def test_refine_index_pure_no_pdf_words(monkeypatch):
    """With an empty PDF (no words on any page), every text element falls back
    to shape source — no crash, schema additive."""
    import src.refine_text as rt

    index = {
        "deck": "x.pptx",
        "slide_size": {"w": 13.333, "h": 7.5, "units": "inches"},
        "slides": [{
            "slide_number": 1,
            "elements": [
                {"id": "s1_e0", "type": "body", "text": "hello",
                 "bbox": {"x": 1.0, "y": 1.0, "w": 4.0, "h": 1.0, "units": "inches"},
                 "needs_vision": False},
                {"id": "s1_e1", "type": "image", "text": "",
                 "bbox": {"x": 0.0, "y": 0.0, "w": 2.0, "h": 2.0, "units": "inches"},
                 "needs_vision": True},
            ],
        }],
    }

    # one page, no words -> every refinable text element falls back to shape.
    monkeypatch.setattr(rt, "_iter_pdf_pages",
                        lambda _p: iter([(960.0, 540.0, [])]))
    out = rt.refine_index(index, "ignored.pdf")
    e0 = out["slides"][0]["elements"][0]
    assert e0["bbox_source"] == "shape"          # refinable but no words
    assert e0["line_bboxes"] == []
    assert "bbox_source" not in out["slides"][0]["elements"][1]  # image untouched


@pytest.mark.skipif(SOFFICE_MISSING, reason="soffice not installed")
@pytest.mark.skipif(not INDEX_JSON.exists(), reason="out/index.json not generated")
def test_refine_index_real_deck_tightens_text(tmp_path):
    import json as _json

    from src.refine_text import refine_index
    from src.render import deck_pdf

    index = _json.load(open(INDEX_JSON))
    pdf = deck_pdf(DECK, tmp_path)
    refined = refine_index(index, pdf)

    slide2 = _slide(refined, 2)
    pdf_text = [e for e in slide2["elements"] if e.get("bbox_source") == "pdf_text"]
    assert pdf_text, "expected at least one refined text element on slide 2"

    sw = refined["slide_size"]["w"]
    sh = refined["slide_size"]["h"]
    eps = 0.05

    # At least one element is strictly tighter horizontally than its shape box.
    assert any(e["bbox"]["w"] < e["shape_bbox"]["w"] - eps for e in pdf_text)

    for e in pdf_text:
        assert e["line_bboxes"], f"{e['id']} has no line boxes"
        for ln in e["line_bboxes"]:
            assert ln["x"] >= e["shape_bbox"]["x"] - eps
            assert ln["y"] >= e["shape_bbox"]["y"] - eps
            assert ln["x"] + ln["w"] <= e["shape_bbox"]["x"] + e["shape_bbox"]["w"] + eps
            assert ln["y"] + ln["h"] <= e["shape_bbox"]["y"] + e["shape_bbox"]["h"] + eps
            assert -eps <= ln["x"] and ln["x"] + ln["w"] <= sw + eps
            assert -eps <= ln["y"] and ln["y"] + ln["h"] <= sh + eps


def test_baseline_correct_drops_box_by_height_fraction():
    from src.refine_text import _baseline_correct
    box = {"x": 1.0, "y": 2.0, "w": 3.0, "h": 0.30, "units": "inches"}
    out = _baseline_correct(box, shift=0.22)             # explicit size-proportional drop
    assert out["x"] == 1.0 and out["w"] == 3.0           # x/w unchanged
    assert out["h"] == 0.30                              # height unchanged
    assert out["y"] == 2.066                            # dropped 0.22 * 0.30 = 0.066"
    # default shift is SHIFT_FRAC (0.0 for tight glyph boxes) -> no drop
    assert _baseline_correct(box)["y"] == 2.0
    assert _baseline_correct(None) is None


def test_union_of_boxes_multiline():
    from src.refine_text import _union_of_boxes
    a = {"x": 1.0, "y": 2.0, "w": 3.0, "h": 0.2, "units": "inches"}
    b = {"x": 1.2, "y": 2.4, "w": 2.0, "h": 0.2, "units": "inches"}
    u = _union_of_boxes([a, None, b])
    assert u == {"x": 1.0, "y": 2.0, "w": 3.0, "h": 0.6, "units": "inches"}
    assert _union_of_boxes([None]) is None


def _pt_word(text, x, y, w=0.5, h=0.15):
    """A PDF word dict in points (y from page top), as _iter_pdf_pages yields."""
    return {"text": text, "x0": x * 72, "x1": (x + w) * 72, "top": y * 72, "bottom": (y + h) * 72}


def _refine_one_slide(monkeypatch, elements, pt_words):
    import src.refine_text as rt
    index = {"deck": "x.pptx", "slide_size": {"w": 13.333, "h": 7.5, "units": "inches"},
             "slides": [{"slide_number": 1, "elements": elements}]}
    monkeypatch.setattr(rt, "_iter_pdf_pages", lambda _p: iter([(13.333 * 72, 540.0, pt_words)]))
    return {e["id"]: e for e in rt.refine_index(index, "ignored.pdf")["slides"][0]["elements"]}


def _el(eid, text, x, y, w, h, typ="body"):
    return {"id": eid, "type": typ, "text": text, "needs_vision": False,
            "bbox": {"x": x, "y": y, "w": w, "h": h, "units": "inches"}}


def test_full_slide_label_does_not_absorb_other_shapes_overflow(monkeypatch):
    """Bug A: a 2-word label authored as a full-slide text box used to own every
    stray glyph on the slide (another shape's overflow) and frame half the slide."""
    els = [_el("s1_e0", "THE PROBLEM", 0, 0, 13.333, 7.5, "diagram_label"),
           _el("s1_e1", "Handouts come back as phone photos", 1.0, 3.4, 6.2, 0.2, "bullet")]
    words = [_pt_word("THE", 0.7, 0.48, 0.3, 0.12), _pt_word("PROBLEM", 1.04, 0.48, 0.7, 0.12),
             _pt_word("Handouts", 1.0, 3.42), _pt_word("come", 1.6, 3.42),
             _pt_word("back", 2.2, 3.42), _pt_word("as", 2.8, 3.42),
             # the bullet wraps: its second line renders below its own box
             _pt_word("phone", 1.0, 3.62), _pt_word("photos", 1.6, 3.62),
             _pt_word(",", 6.4, 4.2, 0.05, 0.05), _pt_word("ng", 1.9, 4.5, 0.2)]  # stray glyphs
    out = _refine_one_slide(monkeypatch, els, words)
    label = out["s1_e0"]
    assert label["bbox_source"] == "pdf_text"
    assert label["bbox"]["w"] < 1.2 and label["bbox"]["h"] < 0.2      # just the two words
    bullet = out["s1_e1"]
    assert bullet["bbox_source"] == "pdf_text" and len(bullet["line_bboxes"]) == 2


def test_select_text_words_drops_stray_runs():
    from src.refine_text import select_text_words
    w = lambda i, t: {"i": i, "text": t, "x0": i, "y0": 0, "x1": i + 0.5, "y1": 0.2}
    words = [w(0, "STUDENT"), w(1, "-"), w(2, "HANDOUT"), w(3, "AUTO"), w(4, "-"), w(5, "GRADING"),
             w(9, ","), w(12, "ng"), w(15, "y")]     # 'ng' is a substring but already covered
    kept, cover = select_text_words("STUDENT-HANDOUT AUTO-GRADING", words)
    assert [x["text"] for x in kept] == ["STUDENT", "-", "HANDOUT", "AUTO", "-", "GRADING"]
    assert cover == 1.0


def test_stray_glyph_never_becomes_a_degenerate_box(monkeypatch):
    """Bug B: a long bullet whose only in-box glyph was a lone '.' refined to a
    0.02 x 0.02 in box; it must fall back to the shape box."""
    els = [_el("s22_e5", "Observation: The device is occasionally missing.", 0.74, 1.97, 12.0, 0.27,
               "bullet")]
    out = _refine_one_slide(monkeypatch, els, [_pt_word(".", 10.97, 1.97, 0.02, 0.02)])
    e = out["s22_e5"]
    assert e["bbox_source"] == "shape" and e["bbox"] == e["shape_bbox"]


def test_text_rendered_below_its_parsed_box_is_found_by_stream_alignment(monkeypatch):
    """Parsed paragraph boxes drift from the render (code listings): each bullet's
    glyphs sit in the NEXT bullet's box. Alignment follows the text, not the box."""
    els = [_el("a", "reg = <0x34>;", 0.74, 1.65, 12.0, 0.22, "bullet"),
           _el("b", "compatible = ti;", 0.74, 1.92, 12.0, 0.22, "bullet")]
    words = [_pt_word("reg", 1.25, 1.9, 0.3), _pt_word("=", 1.6, 1.9, 0.1),
             _pt_word("<0x34>;", 1.8, 1.9, 0.6),
             _pt_word("compatible", 1.25, 2.3, 0.9), _pt_word("=", 2.2, 2.3, 0.1),
             _pt_word("ti;", 2.4, 2.3, 0.2)]
    out = _refine_one_slide(monkeypatch, els, words)
    assert out["a"]["bbox"]["y"] == pytest.approx(1.9) and out["a"]["bbox"]["w"] == pytest.approx(1.15)
    assert out["b"]["bbox"]["y"] == pytest.approx(2.3)


def test_text_whose_math_is_blank_in_the_pdf_still_refines(monkeypatch):
    """Equation boxes: soffice draws the words but leaves the math blank — the
    glyph run reads "Let be languages" for "Let A, B be languages". Its in-order
    matching blocks (let + belanguages) cover every glyph, so the run is the
    element's text with a gap, not a stray."""
    from src.refine_text import select_text_words

    els = [_el("eq", "Let A, B be languages", 1.0, 2.0, 8.0, 1.0)]
    words = [_pt_word("Let", 1.1, 2.1, 0.3), _pt_word("be", 2.0, 2.1, 0.2),
             _pt_word("languages", 2.3, 2.1, 0.9)]
    e = _refine_one_slide(monkeypatch, els, words)["eq"]
    assert e["bbox_source"] == "pdf_text"
    assert len(e["line_bboxes"]) == 1, e["line_bboxes"]
    assert e["bbox"]["x"] == pytest.approx(1.1) and e["bbox"]["x"] + e["bbox"]["w"] == pytest.approx(3.2)
    assert e["bbox"]["y"] == pytest.approx(2.1)
    # the blocks rule itself: the gap run is kept, an unrelated run (another
    # shape's overflow into the box, its own stream run) is still a stray
    w = lambda i, t: {"i": i, "text": t, "x0": i, "y0": 0, "x1": i + 0.5, "y1": 0.2}
    kept, cover = select_text_words("Let A, B be languages",
                                    [w(0, "Let"), w(1, "be"), w(2, "languages"),
                                     w(4, "Quarterly"), w(5, "revenue"), w(6, "grew")])
    assert [x["text"] for x in kept] == ["Let", "be", "languages"]
    assert cover == pytest.approx(14 / 16)


def test_text_with_blank_math_is_located_by_the_stream_when_its_box_drifted(monkeypatch):
    """MIT slide 8: the parser's guessed line sits 0.5 in above the glyphs AND the
    math is blank in the PDF, so neither the exact stream search nor containment
    finds the line. The surviving segments abut in the stream (blank math leaves
    no glyphs), so the run is located by its in-order matching blocks; the
    neighbouring plain line, found exactly, is untouched."""
    els = [_el("math", "The empty language ø is the set with no strings", 0.5, 3.135, 9.0, 0.333),
           _el("plain", "Recognizing languages", 0.5, 3.636, 9.0, 0.4)]
    words = [_pt_word("The", 1.0, 3.94, 0.3), _pt_word("empty", 1.4, 3.94, 0.5),
             _pt_word("language", 2.0, 3.94, 0.8), _pt_word("is", 3.3, 3.94, 0.2),
             _pt_word("the", 3.6, 3.94, 0.3), _pt_word("set", 4.0, 3.94, 0.3),
             _pt_word("with", 4.4, 3.94, 0.4), _pt_word("no", 4.9, 3.94, 0.2),
             _pt_word("strings", 5.2, 3.94, 0.6),
             _pt_word("Recognizing", 1.0, 3.75, 1.0), _pt_word("languages", 2.1, 3.75, 0.9)]
    out = _refine_one_slide(monkeypatch, els, words)
    assert out["math"]["bbox_source"] == "pdf_text"
    assert out["math"]["bbox"]["y"] == pytest.approx(3.94)
    assert out["math"]["bbox"]["x"] == pytest.approx(1.0) and out["math"]["bbox"]["w"] == pytest.approx(4.8)
    assert out["plain"]["bbox"]["y"] == pytest.approx(3.75)


def test_gap_span_with_repeated_anchor_takes_the_occurrence_at_its_own_box(monkeypatch):
    """'is regular' is drawn twice: in B's line above and in A's own line (A's
    math blank: "Then  is regular"). Every occurrence of the anchor is a
    candidate and proximity chooses, so A takes its own line and B keeps hers."""
    els = [_el("b", "This is regular.", 1.0, 2.0, 5.0, 0.3),
           _el("a", "Then L(A) ∪ L(B) is regular", 1.0, 2.5, 5.0, 0.3)]
    words = [_pt_word("This", 1.0, 2.05, 0.4), _pt_word("is", 1.5, 2.05, 0.2),
             _pt_word("regular.", 1.8, 2.05, 0.7),
             _pt_word("Then", 1.0, 2.55, 0.4), _pt_word("is", 2.6, 2.55, 0.2),
             _pt_word("regular", 2.9, 2.55, 0.7)]
    out = _refine_one_slide(monkeypatch, els, words)
    assert out["a"]["bbox_source"] == "pdf_text" and out["a"]["bbox"]["y"] == pytest.approx(2.55)
    assert out["a"]["bbox"]["x"] == pytest.approx(1.0) and out["a"]["bbox"]["w"] == pytest.approx(2.6)
    assert out["b"]["bbox_source"] == "pdf_text" and out["b"]["bbox"]["y"] == pytest.approx(2.05)


def test_gap_span_keeps_its_first_word_despite_an_earlier_copy_of_it(monkeypatch):
    """The title "Let's begin" above the equation line also starts with 'let'. A
    run grows from its anchor by the segments that ABUT it, so the equation line
    keeps its own 'Let' rather than losing it to the earlier copy."""
    els = [_el("t", "Let's begin", 1.0, 0.5, 3.0, 0.4),
           _el("eq", "Let A, B be languages", 1.0, 2.0, 8.0, 1.0)]
    words = [_pt_word("Let's", 1.0, 0.55, 0.5), _pt_word("begin", 1.6, 0.55, 0.5),
             _pt_word("Let", 1.1, 2.1, 0.3), _pt_word("be", 2.0, 2.1, 0.2),
             _pt_word("languages", 2.3, 2.1, 0.9)]
    out = _refine_one_slide(monkeypatch, els, words)
    assert out["eq"]["bbox_source"] == "pdf_text"
    assert out["eq"]["bbox"]["x"] == pytest.approx(1.1) and out["eq"]["bbox"]["w"] == pytest.approx(2.1)
    assert out["t"]["bbox_source"] == "pdf_text" and out["t"]["bbox"]["y"] == pytest.approx(0.55)


def test_gap_span_that_is_another_elements_whole_text_is_left_to_it(monkeypatch):
    """"Theorem 2.1" has no glyphs at all; the run "Theorem" 1 in below spells
    most of it, but it IS the neighbour's text — hers, not a gap span."""
    els = [_el("a", "Theorem 2.1", 1.0, 1.0, 1.5, 0.3), _el("b", "Theorem", 1.0, 2.0, 1.5, 0.3)]
    out = _refine_one_slide(monkeypatch, els, [_pt_word("Theorem", 1.0, 2.05, 0.8)])
    assert out["a"]["bbox_source"] == "shape"
    assert out["b"]["bbox_source"] == "pdf_text"


def test_rescue_leaves_a_gap_span_another_unplaced_element_spells(monkeypatch):
    """Pass 2: A's math line is far from its box; the only run spelling it sits in
    the box of B, a long paragraph whose text spells the run (its first line)
    without containing it verbatim. The rescue judges ownership by the same
    gap-tolerant rule, so A does not take B's line."""
    els = [_el("a", "Then L(A) ∪ L(B) is regular", 1.0, 0.5, 5.0, 0.3),
           _el("b", "Then L is regular and this remark about it runs on for a while",
               1.0, 6.4, 5.0, 0.6)]
    words = [_pt_word("Then", 1.0, 6.55, 0.4), _pt_word("is", 2.6, 6.55, 0.2),
             _pt_word("regular", 2.9, 6.55, 0.7)]
    out = _refine_one_slide(monkeypatch, els, words)
    assert out["a"]["bbox_source"] == "shape"


def test_page_unique_rescue_does_not_steal_an_equal_length_neighbours_glyphs(monkeypatch):
    """A="Proof" whose glyphs are not in the PDF must not take B="Proof." (same
    normalised length, so A sorts first) from 4.5 in away: B's proximity match
    is placed first, then A has no free occurrence left and keeps its shape box."""
    els = [_el("a", "Proof", 1.0, 0.5, 1.0, 0.3), _el("b", "Proof.", 2.5, 4.5, 1.0, 0.4)]
    out = _refine_one_slide(monkeypatch, els, [_pt_word("Proof.", 2.8, 5.0)])
    assert out["a"]["bbox_source"] == "shape"
    assert out["b"]["bbox_source"] == "pdf_text"


def test_page_unique_rescue_leaves_words_inside_an_unplaced_owners_box(monkeypatch):
    """B's own stream match fails (its last word was never drawn) while the word
    A spells sits inside B's box: the rescue must not take it from B, which then
    claims its whole run by containment."""
    els = [_el("a", "Proof", 1.0, 0.5, 1.0, 0.3),
           _el("b", "Proof of revenue 2024 results", 1.0, 4.9, 3.0, 0.4)]
    words = [_pt_word("Proof", 1.0, 5.0), _pt_word("of", 1.6, 5.0),
             _pt_word("revenue", 2.2, 5.0), _pt_word("2024", 2.8, 5.0)]
    out = _refine_one_slide(monkeypatch, els, words)
    assert out["a"]["bbox_source"] == "shape"
    assert out["b"]["bbox_source"] == "pdf_text"
    assert out["b"]["bbox"]["x"] == pytest.approx(1.0)


@pytest.mark.skipif(SOFFICE_MISSING, reason="soffice not installed")
def test_paragraphs_drifted_far_below_parsed_boxes_still_align(tmp_path):
    """Body-placeholder paragraphs inherit their font size from the master
    (python-pptx sees None, the parser guesses 16 pt lines), so the parser's
    guessed line boxes drift further from the rendered glyphs with every
    paragraph — past any fixed proximity radius (TapHole slide 8: 15 paragraphs
    with inherited sizes + normAutofit 92.5 %, the last three lost). A text that
    occurs once in the page's content stream is located by the stream alone."""
    import pypdfium2 as pdfium
    from pptx import Presentation

    from src.parser import parse_deck
    from src.refine_text import refine_index
    from src.render import deck_pdf

    tokens = ["Alpha", "Bravo", "Charlie", "Delta", "Echo", "Foxtrot", "Golf", "Hotel", "India"]
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content: body inherits 32 pt
    slide.shapes.title.text = "Inherited sizes"
    tf = slide.placeholders[1].text_frame
    for i, tok in enumerate(tokens):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{tok} line {i + 1}"
    deck = tmp_path / "inherited_size.pptx"
    prs.save(deck)

    pdf = deck_pdf(deck, tmp_path)
    index = parse_deck(str(deck))
    doc = pdfium.PdfDocument(str(pdf))
    page = doc[0]
    textpage = page.get_textpage()
    page_h = page.get_size()[1]
    glyph_y = {}
    for tok in tokens:
        hit = textpage.search(tok, match_whole_word=True).get_next()
        assert hit, f"{tok} missing from the PDF text layer"
        glyph_y[tok] = (page_h - textpage.get_charbox(hit[0], loose=False)[3]) / 72
    doc.close()

    paras = {e["text"].split()[0]: e for e in index["slides"][0]["elements"] if e["type"] == "bullet"}
    assert set(paras) == set(tokens)
    last = paras[tokens[-1]]["bbox"]
    # precondition: the parser's guess for the last paragraph sits > 2 in above its glyphs
    assert glyph_y[tokens[-1]] - (last["y"] + last["h"]) > 2.0, "drift too small to exercise the bug"

    refine_index(index, pdf)
    for tok in tokens:
        e = paras[tok]
        assert e["line_bboxes"], f"{tok}: not refined ({e['bbox_source']})"
        assert e["line_bboxes"][0]["y"] == pytest.approx(glyph_y[tok], abs=0.1), tok
