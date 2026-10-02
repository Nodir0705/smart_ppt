"""tools.make_image_decks — image-only ground-truth deck generator."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from tools.make_image_decks import main


def _make_deck(path: Path) -> None:
    prs = Presentation()
    for text in ["Hello World", "Second Slide Title"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        tb.text_frame.text = text
    prs.save(str(path))


def _make_png(path: Path) -> None:
    Image.new("RGB", (320, 240), (90, 130, 180)).save(path, "PNG")


def test_make_image_decks(tmp_path, capsys):
    decks_dir = tmp_path / "decks"
    bench_dir = tmp_path / "bench"
    out_dir = tmp_path / "out"
    decks_dir.mkdir()

    _make_deck(decks_dir / "mydeck.pptx")
    (decks_dir / "README.md").write_text("not a deck")

    # a second deck with no rendered slides -> must be skipped, not error
    _make_deck(decks_dir / "norender.pptx")

    slides_dir = bench_dir / "mydeck" / "default" / "slides"
    slides_dir.mkdir(parents=True)
    _make_png(slides_dir / "slide_001.png")
    _make_png(slides_dir / "slide_002.png")

    rc = main(["--decks", str(decks_dir), "--bench", str(bench_dir), "--out", str(out_dir)])
    assert rc == 0

    out_pptx = out_dir / "mydeck__as_images.pptx"
    assert out_pptx.exists()
    prs = Presentation(str(out_pptx))
    assert len(list(prs.slides)) == 2
    for slide in prs.slides:
        pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pics) == 1
        pic = pics[0]
        assert pic.left == 0
        assert pic.top == 0
        assert pic.width == prs.slide_width
        assert pic.height == prs.slide_height

    truth_path = out_dir / "mydeck__truth.json"
    assert truth_path.exists()
    import json
    truth = json.loads(truth_path.read_text())
    texts = [
        (e.get("text") or "")
        for s in truth["slides"]
        for e in s["elements"]
    ]
    assert any("Hello World" in t for t in texts)
    assert any("Second Slide Title" in t for t in texts)
    assert truth["slide_png"] == {
        "1": str(slides_dir / "slide_001.png"),
        "2": str(slides_dir / "slide_002.png"),
    }

    # norender deck was skipped without error and produced no output files
    assert not (out_dir / "norender__as_images.pptx").exists()
    assert not (out_dir / "norender__truth.json").exists()

    out = capsys.readouterr().out
    assert "norender: skip" in out
    assert "mydeck: 2 slides" in out
