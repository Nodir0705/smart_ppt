"""Build image-only ground-truth decks for the OCR / layout / vision comparison.

For each real deck, paste its already-rendered slide PNGs (one per slide, full
bleed) into a fresh pptx — so the slide becomes a single picture with no live
text. The original deck's XML parse (`src.parser.parse_deck`) is then exact
ground truth (text + bbox, same inch coordinate space as the PNG) for what an
OCR/VLM pipeline should recover from the pixels alone.

    python3 -m tools.make_image_decks
    python3 -m tools.make_image_decks --decks decks --bench out/bench --out out/image_decks

Does NOT render slides itself (reads out/bench/<name>/default/slides/*.png,
already produced by tools.deck_bench) and does NOT touch out/bench/ otherwise.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation

from src.parser import parse_deck

_BLANK = 6  # "Blank" layout index in the default template


def _slides_dir(bench_dir: Path, name: str) -> Path:
    return bench_dir / name / "default" / "slides"


def build_one(deck_path: Path, bench_dir: Path, out_dir: Path) -> str | None:
    """Build `<name>__as_images.pptx` + `<name>__truth.json` for one deck.

    Returns the printed summary line, or None if the deck was skipped
    (no rendered slides dir / no PNGs in it).
    """
    name = deck_path.stem
    slides_dir = _slides_dir(bench_dir, name)
    if not slides_dir.is_dir():
        msg = f"{name}: skip (no rendered slides at {slides_dir})"
        print(msg)
        return None
    pngs = sorted(slides_dir.glob("slide_*.png"))
    if not pngs:
        msg = f"{name}: skip (no slide PNGs in {slides_dir})"
        print(msg)
        return None

    orig = Presentation(str(deck_path))
    slide_w, slide_h = orig.slide_width, orig.slide_height

    prs = Presentation()
    prs.slide_width = slide_w
    prs.slide_height = slide_h
    blank = prs.slide_layouts[_BLANK]

    slide_png: dict[str, str] = {}
    for png in pngs:
        n = int(png.stem.rsplit("_", 1)[-1])
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, width=slide_w, height=slide_h)
        slide_png[str(n)] = str(png)

    out_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = out_dir / f"{name}__as_images.pptx"
    prs.save(str(pptx_path))

    truth = parse_deck(str(deck_path), ocr=False)
    truth["slide_png"] = slide_png
    (out_dir / f"{name}__truth.json").write_text(json.dumps(truth, indent=2))

    msg = f"{name}: {len(pngs)} slides -> {pptx_path}"
    print(msg)
    return msg


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description="Build image-only ground-truth decks for OCR/VLM comparison."
    )
    ap.add_argument("--decks", default="decks", help="source decks dir (default: decks)")
    ap.add_argument("--bench", default="out/bench", help="rendered-slides dir (default: out/bench)")
    ap.add_argument("--out", default="out/image_decks", help="output dir (default: out/image_decks)")
    args = ap.parse_args(argv)

    decks_dir = Path(args.decks)
    bench_dir = Path(args.bench)
    out_dir = Path(args.out)

    for deck_path in sorted(decks_dir.glob("*.pptx")):
        build_one(deck_path, bench_dir, out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
