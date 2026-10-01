"""ocr_bakeoff.py — static OCR bake-off: run engines over slide PNGs + icon crops.

Produces `out/bakeoff/raw/<engine>.json` in the standard raw format (one entry
per detected line/word, bbox in slide-space inches), scored separately by
tools.bakeoff_score. Never touches src/ocr.py, src/parser.py, src/build_index.py
or out/bench — all output goes to out/bakeoff/.

Engines:
  tesseract-current  — exact src/ocr.py settings (psm 3, conf>=60, alpha>=3)
  tesseract-raw      — same psm, no filter (shows what the filter costs)
  paddleocr-mobile/-server — via a subprocess in ~/.venvs/smartppt-paddle
  (paddleocr-server needs considerably more RAM than paddleocr-mobile; it
  OOM-killed on an 8-9GB-available shared host during this bake-off — pass
  it explicitly via --engines once you know the box can hold it)

    python3 -m tools.ocr_bakeoff                       # default: the 3 engines that ran clean here
    python3 -m tools.ocr_bakeoff --engines paddleocr-server
    python3 -m tools.ocr_bakeoff --determinism-decks test_deck   # rerun + diff
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path

from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
TRUTH_DIR = Path("/home/netvision/smartppt/out/image_decks")
SMARTPPT_ROOT = Path("/home/netvision/smartppt")
OUT_DIR = Path("/home/netvision/smartppt/out/bakeoff")
RAW_DIR = OUT_DIR / "raw"
DECK_NAMES = [
    "test_deck",
    "AI_Handout_Checker",
    "Open-Q 845 Board Technical Review – Android 10 System Architecture",
]

PADDLE_PY = Path.home() / ".venvs" / "smartppt-paddle" / "bin" / "python"
PADDLE_WORKER = ROOT / "tools" / "_paddle_worker.py"

# reproduce src/ocr.py's filter exactly (see that module's docstring/constants)
DEFAULT_PSM = 3
MIN_CONF = 60.0
MIN_ALPHA_CHARS = 3
_TESS = "tesseract"
_ALPHA_RE = re.compile(r"[^A-Za-z]")
ICON_MAX_AREA_SQIN = 0.5
ICON_PAD_PX = 2


def _keep(token: str) -> bool:
    return len(_ALPHA_RE.sub("", token)) >= MIN_ALPHA_CHARS


# ---------- truth / IO helpers ----------

def load_truth(name: str) -> dict:
    return json.loads((TRUTH_DIR / f"{name}__truth.json").read_text())


def png_path_for(truth: dict, slide_no) -> Path:
    rel = truth["slide_png"][str(slide_no)]
    return SMARTPPT_ROOT / rel


# ---------- tesseract engine ----------

def _tesseract_tsv(png_path: Path, psm: int = DEFAULT_PSM) -> list[dict]:
    r = subprocess.run(
        [_TESS, str(png_path), "stdout", "--psm", str(psm), "-l", "eng", "tsv"],
        capture_output=True, text=True, timeout=120, check=False,
    )
    if r.returncode != 0:
        return []
    return list(csv.DictReader(r.stdout.splitlines(), delimiter="\t"))


def _group_lines(rows: list[dict], png_w: int, png_h: int,
                  slide_w: float, slide_h: float, filtered: bool) -> list[dict]:
    """Group TSV level-5 word rows into line entries; convert px -> slide inches."""
    lines: dict[tuple, dict] = {}
    order: list[tuple] = []
    for row in rows:
        if row.get("level") != "5":
            continue
        try:
            conf = float(row.get("conf", "-1"))
        except (TypeError, ValueError):
            continue
        token = (row.get("text") or "").strip()
        if not token:
            continue
        if filtered and not (conf >= MIN_CONF and _keep(token)):
            continue
        key = (row["block_num"], row["par_num"], row["line_num"])
        x, y, w, h = int(row["left"]), int(row["top"]), int(row["width"]), int(row["height"])
        if key not in lines:
            lines[key] = {"words": [], "x0": x, "y0": y, "x1": x + w, "y1": y + h, "confs": []}
            order.append(key)
        line = lines[key]
        line["words"].append(token)
        line["x0"] = min(line["x0"], x)
        line["y0"] = min(line["y0"], y)
        line["x1"] = max(line["x1"], x + w)
        line["y1"] = max(line["y1"], y + h)
        line["confs"].append(conf)
    out = []
    sx, sy = slide_w / png_w, slide_h / png_h
    for key in order:
        line = lines[key]
        text = " ".join(line["words"])
        bbox = [line["x0"] * sx, line["y0"] * sy,
                (line["x1"] - line["x0"]) * sx, (line["y1"] - line["y0"]) * sy]
        conf = sum(line["confs"]) / len(line["confs"])
        out.append({"text": text, "bbox": bbox, "conf": conf})
    return out


def run_tesseract_deck(truth: dict, filtered: bool) -> tuple[dict, float]:
    slide_w, slide_h = truth["slide_size"]["w"], truth["slide_size"]["h"]
    per_slide: dict[str, list] = {}
    t0 = time.perf_counter()
    n = 0
    for s in truth["slides"]:
        no = s["slide_number"]
        png = png_path_for(truth, no)
        if not png.exists():
            continue
        with Image.open(png) as im:
            png_w, png_h = im.size
        rows = _tesseract_tsv(png)
        per_slide[str(no)] = _group_lines(rows, png_w, png_h, slide_w, slide_h, filtered)
        n += 1
    elapsed = time.perf_counter() - t0
    return per_slide, (elapsed / n if n else 0.0)


def run_tesseract_icon(png_path: Path, filtered: bool) -> list[dict]:
    rows = _tesseract_tsv(png_path)
    with Image.open(png_path) as im:
        w, h = im.size
    # inches irrelevant for icon FP check; keep pixel-space bbox for debugging
    return _group_lines(rows, w, h, w, h, filtered)


# ---------- icon / negative crops ----------

def icon_crops(deck_name: str, truth: dict) -> list[tuple[str, Image.Image]]:
    slide_w, slide_h = truth["slide_size"]["w"], truth["slide_size"]["h"]
    out = []
    for s in truth["slides"]:
        no = s["slide_number"]
        png = png_path_for(truth, no)
        if not png.exists():
            continue
        im = Image.open(png).convert("RGB")
        png_w, png_h = im.size
        for el in s["elements"]:
            if el.get("type") != "image":
                continue
            b = el["bbox"]
            if b["w"] * b["h"] >= ICON_MAX_AREA_SQIN:
                continue
            x0 = b["x"] / slide_w * png_w
            y0 = b["y"] / slide_h * png_h
            x1 = (b["x"] + b["w"]) / slide_w * png_w
            y1 = (b["y"] + b["h"]) / slide_h * png_h
            x0 = max(0, int(x0) - ICON_PAD_PX)
            y0 = max(0, int(y0) - ICON_PAD_PX)
            x1 = min(png_w, int(x1) + ICON_PAD_PX)
            y1 = min(png_h, int(y1) + ICON_PAD_PX)
            if x1 <= x0 or y1 <= y0:
                continue
            label = f"{deck_name}/{no}/{el['id']}"
            out.append((label, im.crop((x0, y0, x1, y1))))
    return out


def image_no_text_crop() -> tuple[str, Image.Image]:
    """The corpus's image_no_text deck: shapes, no text, no alt — OCR must not
    hallucinate a pointable. Built via tools.make_corpus into a tmp dir."""
    from pptx import Presentation
    from tools.make_corpus import build_corpus

    with tempfile.TemporaryDirectory(prefix="bakeoff_corpus_") as td:
        decks = build_corpus(td)
        spec = next(d for d in decks if d.name == "image_no_text")
        prs = Presentation(spec.path)
        pic = next(
            shp for slide in prs.slides for shp in slide.shapes
            if shp.shape_type == 13  # MSO_SHAPE_TYPE.PICTURE
        )
        blob = pic.image.blob
        img = Image.open(io.BytesIO(blob)).convert("RGB")
        img.load()  # materialize before the tmpdir goes away
        return "corpus/image_no_text", img


# ---------- paddleocr engine (subprocess in the dedicated venv) ----------

def run_paddle_deck(truth: dict, deck_name: str, model: str, tmp_dir: Path) -> tuple[dict, float, dict]:
    """Batch every slide of `deck` through one worker invocation (model loads once).
    Returns (per_slide, seconds_per_slide, extra_meta)."""
    spec = []
    slide_w, slide_h = truth["slide_size"]["w"], truth["slide_size"]["h"]
    for s in truth["slides"]:
        no = s["slide_number"]
        png = png_path_for(truth, no)
        if png.exists():
            spec.append({"key": str(no), "image": str(png)})
    spec_path = tmp_dir / f"spec_{model}_{abs(hash(deck_name))}.json"
    out_path = tmp_dir / f"out_{model}_{abs(hash(deck_name))}.json"
    spec_path.write_text(json.dumps(spec))
    _run_paddle_worker(model, spec_path, out_path)
    raw = json.loads(out_path.read_text())
    per_slide = {}
    for no_str, dets in raw["results"].items():
        entries = []
        for d in dets:
            xs = [p[0] for p in d["poly"]]
            ys = [p[1] for p in d["poly"]]
            x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
            # poly is in source-image pixels; need that slide's png size
            png_w, png_h = d["img_w"], d["img_h"]
            sx, sy = slide_w / png_w, slide_h / png_h
            entries.append({
                "text": d["text"],
                "bbox": [x0 * sx, y0 * sy, (x1 - x0) * sx, (y1 - y0) * sy],
                "conf": d["conf"],
            })
        per_slide[no_str] = entries
    sec_per_slide = raw["seconds_total"] / max(1, len(spec))
    meta = {"model_info": raw.get("model_info")}
    return per_slide, sec_per_slide, meta


def run_paddle_icons(crops: list[tuple[str, Image.Image]], model: str, tmp_dir: Path) -> dict:
    spec = []
    crop_dir = tmp_dir / "icons"
    crop_dir.mkdir(exist_ok=True)
    for label, img in crops:
        safe = re.sub(r"[^A-Za-z0-9]+", "_", label)
        p = crop_dir / f"{safe}.png"
        img.save(p)
        spec.append({"key": label, "image": str(p)})
    spec_path = tmp_dir / f"spec_icons_{model}.json"
    out_path = tmp_dir / f"out_icons_{model}.json"
    spec_path.write_text(json.dumps(spec))
    _run_paddle_worker(model, spec_path, out_path)
    raw = json.loads(out_path.read_text())
    return {k: [{"text": d["text"], "conf": d["conf"]} for d in v]
            for k, v in raw["results"].items()}


def _run_paddle_worker(model: str, spec_path: Path, out_path: Path) -> None:
    if not PADDLE_PY.exists():
        raise RuntimeError(f"paddle venv python not found at {PADDLE_PY}; run setup first")
    r = subprocess.run(
        [str(PADDLE_PY), str(PADDLE_WORKER), "--model", model,
         "--spec", str(spec_path), "--out", str(out_path)],
        capture_output=True, text=True, timeout=3600,
    )
    if r.returncode != 0:
        raise RuntimeError(f"paddle worker failed ({model}):\n{r.stdout}\n{r.stderr}")


# ---------- orchestration ----------

def build_icon_set(deck_truths: dict[str, dict]) -> list[tuple[str, Image.Image]]:
    crops = []
    for name, truth in deck_truths.items():
        crops.extend(icon_crops(name, truth))
    crops.append(image_no_text_crop())
    return crops


def write_tesseract_engine(engine: str, filtered: bool, deck_truths: dict,
                            icons: list[tuple[str, Image.Image]]) -> None:
    out: dict = {"_unit": "line"}
    sec_sum, sec_n = 0.0, 0
    for name, truth in deck_truths.items():
        per_slide, sec_per_slide = run_tesseract_deck(truth, filtered)
        out[name] = per_slide
        if per_slide:
            sec_sum += sec_per_slide
            sec_n += 1
    out["_seconds_per_slide"] = sec_sum / sec_n if sec_n else 0.0

    icon_results = {}
    with tempfile.TemporaryDirectory(prefix="bakeoff_icons_") as td:
        tdp = Path(td)
        for label, img in icons:
            p = tdp / (re.sub(r"[^A-Za-z0-9]+", "_", label) + ".png")
            img.save(p)
            entries = run_tesseract_icon(p, filtered)
            icon_results[label] = [{"text": e["text"], "conf": e["conf"]} for e in entries]
    out["_icons"] = icon_results
    out["_deterministic"] = None  # filled in separately by --determinism-decks

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    (RAW_DIR / f"{engine}.json").write_text(json.dumps(out, ensure_ascii=False))
    print(f"wrote {RAW_DIR / (engine + '.json')}")


def write_paddle_engine(engine: str, model: str, deck_truths: dict,
                         icons: list[tuple[str, Image.Image]]) -> None:
    out: dict = {"_unit": "line"}
    sec_sum, sec_n = 0.0, 0
    model_info = None
    with tempfile.TemporaryDirectory(prefix="bakeoff_paddle_") as td:
        tdp = Path(td)
        for name, truth in deck_truths.items():
            per_slide, sec_per_slide, meta = run_paddle_deck(truth, name, model, tdp)
            out[name] = per_slide
            model_info = meta.get("model_info") or model_info
            if per_slide:
                sec_sum += sec_per_slide
                sec_n += 1
        out["_seconds_per_slide"] = sec_sum / sec_n if sec_n else 0.0
        out["_model_info"] = model_info
        out["_icons"] = run_paddle_icons(icons, model, tdp)
    out["_deterministic"] = None

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    (RAW_DIR / f"{engine}.json").write_text(json.dumps(out, ensure_ascii=False))
    print(f"wrote {RAW_DIR / (engine + '.json')}")


def check_determinism(engine: str, deck_names: list[str], deck_truths: dict) -> bool:
    """Rerun `engine` on the given decks and diff against the already-written
    raw file (deck entries only, ignoring the _seconds_per_slide/_icons sidecars
    which are expected to vary in timing/order of collection)."""
    path = RAW_DIR / f"{engine}.json"
    first = json.loads(path.read_text())
    if engine == "tesseract-current":
        second = {name: run_tesseract_deck(deck_truths[name], True)[0] for name in deck_names}
    elif engine == "tesseract-raw":
        second = {name: run_tesseract_deck(deck_truths[name], False)[0] for name in deck_names}
    else:
        model = {"paddleocr": "mobile", "paddleocr-mobile": "mobile",
                  "paddleocr-server": "server"}[engine]
        with tempfile.TemporaryDirectory(prefix="bakeoff_det_") as td:
            second = {name: run_paddle_deck(deck_truths[name], name, model, Path(td))[0]
                      for name in deck_names}
    identical = all(first.get(name) == second[name] for name in deck_names)
    first["_deterministic"] = identical
    path.write_text(json.dumps(first, ensure_ascii=False))
    return identical


ENGINES_TESS = {"tesseract-current": True, "tesseract-raw": False}
ENGINES_PADDLE = {"paddleocr-mobile": "mobile", "paddleocr-server": "server"}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--engines", default="tesseract-current,tesseract-raw,paddleocr-mobile",
                    help="comma-separated engine names (paddleocr-server excluded by default; it OOMs on low-RAM hosts)")
    ap.add_argument("--determinism-decks", default="",
                    help="comma-separated deck names to rerun for the determinism check (empty = skip)")
    args = ap.parse_args()

    deck_truths = {name: load_truth(name) for name in DECK_NAMES}
    icons = build_icon_set(deck_truths)
    print(f"icon/negative crops: {len(icons)}")

    engines = [e.strip() for e in args.engines.split(",") if e.strip()]
    for engine in engines:
        if engine in ENGINES_TESS:
            write_tesseract_engine(engine, ENGINES_TESS[engine], deck_truths, icons)
        elif engine in ENGINES_PADDLE:
            write_paddle_engine(engine, ENGINES_PADDLE[engine], deck_truths, icons)
        else:
            print(f"unknown engine {engine!r}; skipping", file=sys.stderr)

    det_decks = [d.strip() for d in args.determinism_decks.split(",") if d.strip()]
    for engine in engines:
        if det_decks:
            ok = check_determinism(engine, det_decks, deck_truths)
            print(f"determinism[{engine}] on {det_decks}: {ok}")


if __name__ == "__main__":
    main()
